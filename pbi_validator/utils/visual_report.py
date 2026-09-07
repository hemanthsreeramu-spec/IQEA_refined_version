"""
Reconciliation and reporting for the Visual Comparison page.

Merges Engine 1 (utils.visual_diff) and Engine 2 (utils.visual_diff_llm)
findings into one grid with an agreement flag, and renders the detailed Excel
workbook.

Division of authority: the Python engine is authoritative for geometry and
pixel-level facts; the LLM engine is authoritative for semantics (what a text
change means, whether a chart type really changed). Agreement between the two
raises confidence; a finding from only one engine is still reported, labelled
so the reader knows how much weight to give it.
"""
from __future__ import annotations

import re
from difflib import SequenceMatcher
from io import BytesIO

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SEVERITY_RANK = {"info": 0, "minor": 1, "major": 2, "critical": 3}

# Categories that describe the same underlying defect across engines.
_EQUIVALENT = (
    {"visual_content_change", "chart_type_change", "color_theme_change"},
    {"position_change", "slicer_position_change"},
    {"title_text_change", "truncation", "legend_axis_change"},
    {"value_change"},
    {"missing_visual", "visual_count_change"},
    {"extra_visual", "visual_count_change"},
    {"render_error"},
    {"slicer_selection_change"},
    {"size_change"},
)

FINDING_COLUMNS = ["Engine", "Agreement", "Severity", "Visual", "Visual Name",
                   "Category", "Check", "Source State", "Target State",
                   "Metric / Evidence", "Description", "Recommendation"]


def _norm_name(s: str) -> str:
    return re.sub(r"[^a-z0-9 ]", "", str(s or "").lower()).strip()


def _same_visual(a: str, b: str) -> bool:
    na, nb = _norm_name(a), _norm_name(b)
    if not na or not nb:
        return False
    if na == nb or na in nb or nb in na:
        return True
    return SequenceMatcher(None, na, nb).ratio() > 0.75


def _categories_align(a: str, b: str) -> bool:
    if a == b:
        return True
    return any(a in grp and b in grp for grp in _EQUIVALENT)


def merge_findings(py_res: dict | None, llm_res: dict | None) -> list[dict]:
    """
    Combine both engines' findings, tagging each with an agreement flag:
    'both' (corroborated), 'python_only', or 'llm_only'.
    """
    py_f = [dict(f, engine="python") for f in ((py_res or {}).get("findings") or [])]
    llm_f = [dict(f, engine="llm") for f in ((llm_res or {}).get("findings") or [])]

    if not py_f:
        for f in llm_f:
            f["agreement"] = "llm_only" if py_res else "llm"
        return llm_f
    if not llm_f:
        for f in py_f:
            f["agreement"] = "python_only" if llm_res else "python"
        return py_f

    matched_llm: set[int] = set()
    for pf in py_f:
        pf["agreement"] = "python_only"
        for j, lf in enumerate(llm_f):
            if j in matched_llm:
                continue
            if (_categories_align(pf["category"], lf["category"])
                    and _same_visual(pf.get("visual_name", ""),
                                     lf.get("visual_name", ""))):
                pf["agreement"] = "both"
                lf["agreement"] = "both"
                matched_llm.add(j)
                break
    for j, lf in enumerate(llm_f):
        lf.setdefault("agreement", "both" if j in matched_llm else "llm_only")

    combined = py_f + llm_f
    combined.sort(key=lambda f: (-SEVERITY_RANK.get(f["severity"], 0),
                                 0 if f["agreement"] == "both" else 1,
                                 str(f.get("visual_id", ""))))
    return combined


def combined_verdict(findings: list[dict], fail_on_major: bool = True) -> str:
    worst = max((SEVERITY_RANK.get(f["severity"], 0) for f in findings), default=0)
    fail_at = 2 if fail_on_major else 3
    return ("FAIL" if worst >= fail_at
            else "PASS_WITH_WARNINGS" if worst >= 1 else "PASS")


def findings_df(findings: list[dict]) -> pd.DataFrame:
    rows = [{
        "Engine": {"python": "Python", "llm": "LLM"}.get(f.get("engine"), f.get("engine")),
        "Agreement": {"both": "✅ Both engines",
                      "python_only": "Python only",
                      "llm_only": "LLM only",
                      "python": "Python",
                      "llm": "LLM"}.get(f.get("agreement", ""), ""),
        "Severity": f.get("severity", ""),
        "Visual": f.get("visual_id", ""),
        "Visual Name": f.get("visual_name", "") or "(untitled)",
        "Category": f.get("category", ""),
        "Check": f.get("check", ""),
        "Source State": f.get("source_state", ""),
        "Target State": f.get("target_state", ""),
        "Metric / Evidence": f.get("metric", ""),
        "Description": f.get("description", ""),
        "Recommendation": f.get("recommendation", ""),
    } for f in findings]
    return pd.DataFrame(rows, columns=FINDING_COLUMNS)


# ── Excel ─────────────────────────────────────────────────────────────────────
_HDR_FILL = PatternFill("solid", fgColor="E8650A")
_HDR_FONT = Font(bold=True, color="FFFFFF", size=11)
_SEV_FILL = {
    "critical": PatternFill("solid", fgColor="FFC7CE"),
    "major":    PatternFill("solid", fgColor="FFD9B3"),
    "minor":    PatternFill("solid", fgColor="FFEB9C"),
    "info":     PatternFill("solid", fgColor="DDEBF7"),
}
_VERDICT_FILL = {
    "PASS":               PatternFill("solid", fgColor="C6EFCE"),
    "PASS_WITH_WARNINGS": PatternFill("solid", fgColor="FFEB9C"),
    "WARN":               PatternFill("solid", fgColor="FFEB9C"),
    "FAIL":               PatternFill("solid", fgColor="FFC7CE"),
}
_THIN = Side(style="thin")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_WRAP_TOP = Alignment(wrap_text=True, vertical="top")
_WRAP_CTR = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _write_df(ws, df: pd.DataFrame, widths: dict | None = None,
              fill_col: str | None = None, fill_map: dict | None = None):
    for c, name in enumerate(df.columns, 1):
        cell = ws.cell(row=1, column=c, value=str(name))
        cell.fill, cell.font = _HDR_FILL, _HDR_FONT
        cell.alignment, cell.border = _WRAP_CTR, _BORDER

    for r, row in enumerate(df.itertuples(index=False), 2):
        key = ""
        if fill_col and fill_col in df.columns:
            key = str(row[list(df.columns).index(fill_col)]).lower()
        fill = (fill_map or {}).get(key)
        for c, val in enumerate(row, 1):
            cell = ws.cell(row=r, column=c,
                           value="" if val is None else str(val))
            cell.alignment, cell.border = _WRAP_TOP, _BORDER
            if fill:
                cell.fill = fill

    for c, name in enumerate(df.columns, 1):
        letter = get_column_letter(c)
        ws.column_dimensions[letter].width = (widths or {}).get(str(name), 18)
    ws.freeze_panes = "A2"


def export_visual_comparison_excel(combined: list[dict],
                                   py_res: dict | None,
                                   llm_res: dict | None,
                                   verdict: str,
                                   meta: dict | None = None) -> bytes:
    """
    Detailed workbook:
      Summary              verdict, counts, page metrics from both engines
      Findings             the merged grid, colour-coded by severity
      Per-Visual Metrics   every Python-engine measurement per visual
      LLM Inventory        what the LLM saw in each screenshot (two_pass only)
    """
    wb = openpyxl.Workbook()

    # ── Summary ───────────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Summary"
    rows: list[tuple] = [("PBI Visual Comparison Report", "")]
    for k, v in (meta or {}).items():
        rows.append((k, v))
    rows.append(("", ""))
    rows.append(("OVERALL VERDICT", verdict))

    sev_counts = {k: 0 for k in SEVERITY_RANK}
    for f in combined:
        sev_counts[f.get("severity", "info")] = \
            sev_counts.get(f.get("severity", "info"), 0) + 1
    rows.append(("Total findings", len(combined)))
    for s in ("critical", "major", "minor", "info"):
        rows.append((f"  {s} findings", sev_counts.get(s, 0)))
    agree = sum(1 for f in combined if f.get("agreement") == "both")
    rows.append(("Corroborated by both engines", agree))

    for label, res in (("PYTHON ENGINE", py_res), ("LLM ENGINE", llm_res)):
        if not res:
            continue
        rows.append(("", ""))
        rows.append((label, res.get("verdict", "")))
        if res.get("summary"):
            rows.append(("Summary", res["summary"]))
        for k, v in (res.get("page_metrics") or {}).items():
            rows.append((f"  {k}", v))
        for w in (res.get("warnings") or []):
            rows.append(("  ⚠ Warning", w))

    for r, (k, v) in enumerate(rows, 1):
        a = ws.cell(row=r, column=1, value=str(k))
        b = ws.cell(row=r, column=2, value="" if v is None else str(v))
        b.alignment = _WRAP_TOP
        if r == 1:
            a.font = Font(bold=True, size=14, color="1B2A4A")
        elif str(k) in ("OVERALL VERDICT",) or str(k).endswith("ENGINE"):
            a.font = Font(bold=True, size=11)
            b.font = Font(bold=True, size=11)
            fill = _VERDICT_FILL.get(str(v))
            if fill:
                b.fill = fill
        elif not str(k).startswith("  "):
            a.font = Font(bold=True)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 90

    # ── Findings ──────────────────────────────────────────────────────────────
    ws = wb.create_sheet("Findings")
    df = findings_df(combined)
    if df.empty:
        df = pd.DataFrame([{c: "" for c in FINDING_COLUMNS}])
        df.loc[0, "Description"] = "No differences detected."
    _write_df(ws, df, widths={
        "Engine": 9, "Agreement": 16, "Severity": 10, "Visual": 8,
        "Visual Name": 26, "Category": 22, "Check": 26,
        "Source State": 30, "Target State": 30, "Metric / Evidence": 46,
        "Description": 60, "Recommendation": 46,
    }, fill_col="Severity", fill_map=_SEV_FILL)

    # ── Per-Visual Metrics ────────────────────────────────────────────────────
    if py_res and py_res.get("visual_table"):
        ws = wb.create_sheet("Per-Visual Metrics")
        vdf = pd.DataFrame(py_res["visual_table"])
        _write_df(ws, vdf, widths={"Name": 30, "Verdict": 10},
                  fill_col="Verdict", fill_map={k.lower(): v for k, v
                                                in _VERDICT_FILL.items()})

    # ── LLM Inventory ─────────────────────────────────────────────────────────
    if llm_res and llm_res.get("inventory_table"):
        ws = wb.create_sheet("LLM Inventory")
        _write_df(ws, pd.DataFrame(llm_res["inventory_table"]),
                  widths={"Title": 30, "Values": 40, "Selected": 24})

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
